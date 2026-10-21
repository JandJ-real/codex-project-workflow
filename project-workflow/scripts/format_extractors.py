"""Local text views, never execution, for formats beyond PDF and Word.

Specialized readers may register their existing output for any other binary file.
"""
import csv
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
import json
from pathlib import Path
import zipfile


def loc(**where):
    return '<!-- loc ' + json.dumps(where, ensure_ascii=False) + ' -->\n\n'


class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.blocks = []

    def handle_starttag(self, tag, attrs):
        if tag in {'script','style'}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {'script','style'}:
            self.hidden = max(0, self.hidden-1)

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.blocks.append((self.getpos()[0], data.strip()))


def extract_additional(source, extraction_type):
    suffix = source.suffix.lower()
    parts, warnings, gaps, structure = [], [], [], {}
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                for row in sheet:
                    nonempty = [cell for cell in row if cell.value is not None]
                    if nonempty:
                        cell_range = nonempty[0].coordinate + ':' + nonempty[-1].coordinate
                        parts.append(loc(sheet=sheet.title, cell=cell_range, row=nonempty[0].row) +
                                     ' | '.join('' if c.value is None else str(c.value) for c in
                                                row[nonempty[0].column-1:nonempty[-1].column]) + '\n\n')
                    for cell in row:
                        if cell.value is not None:
                            if cell.data_type == 'f':
                                gaps.append(f'{sheet.title}!{cell.coordinate}: formula not evaluated')
            structure['sheets'] = book.sheetnames
        finally:
            book.close()
        warnings.append('Workbook values and formulas only; charts, images, styles and formula results are not verified.')
    elif suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation(source)
        for number, slide in enumerate(deck.slides, 1):
            for shape_no, shape in enumerate(slide.shapes, 1):
                prefix = loc(slide=number, shape=shape_no)
                if shape.has_text_frame and shape.text.strip():
                    parts.append(prefix + shape.text + '\n\n')
                elif shape.has_table:
                    for row_no, row in enumerate(shape.table.rows, 1):
                        parts.append(loc(slide=number, shape=shape_no, row=row_no) +
                                     ' | '.join(c.text for c in row.cells) + '\n\n')
                else:
                    gaps.append(f'slide {number} shape {shape_no}: non-text content')
            if slide.has_notes_slide:
                text = slide.notes_slide.notes_text_frame.text
                if text.strip():
                    parts.append(loc(slide=number, notes=True) + text + '\n\n')
        structure['slides'] = len(deck.slides)
        warnings.append('Slide text, tables and notes only; layout, charts and visual content are not verified.')
    elif suffix == '.ipynb':
        notebook = json.loads(source.read_text(encoding='utf-8-sig'))
        for number, cell in enumerate(notebook.get('cells', []), 1):
            text = cell.get('source', [])
            text = ''.join(text) if isinstance(text, list) else str(text)
            parts.append(loc(cell=number, cell_type=cell.get('cell_type','unknown')) + text + '\n\n')
            if cell.get('outputs'):
                gaps.append(f'cell {number}: outputs omitted')
        warnings.append('Notebook source only; cells are never executed and outputs are not interpreted.')
    elif suffix in {'.html','.htm'}:
        parser = VisibleHTML()
        parser.feed(source.read_text(encoding='utf-8-sig'))
        parts = [loc(line=line) + text + '\n\n' for line, text in parser.blocks]
        warnings.append('Static visible HTML text; scripts, styles, linked resources and layout are not executed or verified.')
    elif suffix in {'.csv','.tsv'}:
        with source.open(encoding='utf-8-sig', newline='') as stream:
            for number, row in enumerate(csv.reader(stream, delimiter='\t' if suffix=='.tsv' else ','), 1):
                parts.append(loc(row=number) + ' | '.join(row) + '\n\n')
        warnings.append('Delimited rows only; values are not computed or interpreted.')
    elif suffix in {'.eml'}:
        mail = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
        for number, part in enumerate(mail.walk(), 1):
            if part.get_content_disposition() == 'attachment':
                gaps.append(f'attachment {part.get_filename()}: not extracted')
            elif part.get_content_type() == 'text/plain':
                parts.append(loc(part=number) + part.get_content() + '\n\n')
        warnings.append('Plain-text message bodies only; attachments, HTML-only parts and remote content are not read.')
    elif suffix in {'.zip'}:
        with zipfile.ZipFile(source) as archive:
            for entry in archive.infolist():
                parts.append(loc(member=entry.filename) + f'{entry.filename} ({entry.file_size} bytes)\n\n')
        gaps.append('Archive manifest only; member contents are not extracted.')
    elif suffix in {'.txt','.md','.rst','.log','.json','.jsonl','.yaml','.yml','.xml'}:
        for number, line in enumerate(source.read_text(encoding='utf-8-sig').splitlines(), 1):
            if line.strip():
                parts.append(loc(line=number) + line + '\n\n')
        warnings.append('Text view only; machine-readable originals remain authoritative for calculations and edits.')
    else:
        raise ValueError(f'No built-in reader for {suffix or "this format"}. Use the applicable format skill once, '
                         'save a source-located Markdown/text view, then register it with --from-markdown. '
                         'Images need OCR/description, media needs a timestamped transcript, archives need a member index. '
                         'Never execute embedded code merely to build a cache.')
    if not parts:
        gaps.append('No supported text extracted')
        parts = ['[No supported text extracted; use a specialized reader.]\n']
    return extraction_type(''.join(parts), 'local-'+suffix.lstrip('.'), len(parts), warnings, gaps, structure)
