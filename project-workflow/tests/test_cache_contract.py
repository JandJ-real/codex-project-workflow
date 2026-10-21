"""Regressions: hidden edits, converter changes, output tampering and format imports."""
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
from document_cache import cache_document


class CacheContractTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.cache = self.root / 'cache'

    def pdf(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / 'rules.pdf'
        c = Canvas(str(path), pageCompression=0)
        c.drawString(40, 720, 'Budget ceiling: 54 hours.')
        c.save()
        return path

    def test_same_size_same_mtime_source_edit_invalidates_cache(self):
        source = self.pdf()
        first = cache_document(source, self.cache)
        stamp = source.stat()
        source.write_bytes(source.read_bytes().replace(b'54 hours', b'86 hours'))
        os.utime(source, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
        self.assertEqual(source.stat().st_size, stamp.st_size)
        result = cache_document(source, self.cache)
        self.assertEqual(result['status'], 'created')
        self.assertIn('86 hours', Path(first['markdown']).read_text(encoding='utf-8'))

    def test_touch_without_content_change_reuses_without_rewriting(self):
        source = self.pdf()
        first = cache_document(source, self.cache)
        out = Path(first['markdown'])
        stamp = out.stat().st_mtime_ns
        os.utime(source, ns=(source.stat().st_atime_ns, source.stat().st_mtime_ns + 2000000000))
        second = cache_document(source, self.cache)
        self.assertEqual(second['status'], 'reused')
        self.assertEqual(out.stat().st_mtime_ns, stamp)

    def test_cache_records_content_signature_and_separate_quality(self):
        result = cache_document(self.pdf(), self.cache)
        record = json.loads(Path(result['metadata']).read_text(encoding='utf-8'))
        self.assertEqual(len(record['source']['sha256']), 64)
        self.assertIn('revision', record['extractor'])
        self.assertIn('dependencies', record['extractor'])
        self.assertEqual(result['quality_info']['semantic_verification'], 'not_performed')
        self.assertEqual(result['quality_info']['visual_verification'], 'not_performed')

    def test_converter_signature_change_rebuilds_unchanged_source(self):
        source = self.pdf()
        first = cache_document(source, self.cache)
        meta = Path(first['metadata'])
        record = json.loads(meta.read_text(encoding='utf-8'))
        record.setdefault('extractor', {})['revision'] = -1
        meta.write_text(json.dumps(record), encoding='utf-8')
        self.assertEqual(cache_document(source, self.cache)['status'], 'created')

    def test_same_stat_markdown_tampering_is_not_reused(self):
        source = self.pdf()
        first = cache_document(source, self.cache)
        md = Path(first['markdown'])
        stamp = md.stat()
        md.write_bytes(md.read_bytes().replace(b'54 hours', b'00 hours'))
        os.utime(md, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
        result = cache_document(source, self.cache)
        self.assertEqual(result['status'], 'created')
        self.assertIn('54 hours', md.read_text(encoding='utf-8'))

    def test_external_extraction_for_complex_binary_is_reusable(self):
        source = self.root / 'scan.tiff'
        source.write_bytes(b'fixture representing an already-read binary')
        extracted = self.root / 'ocr.md'
        extracted.write_text('## Image 1\n\nTracking code: SWIFT-38\n', encoding='utf-8')
        first = cache_document(source, self.cache, from_markdown=extracted)
        self.assertEqual(cache_document(source, self.cache)['status'], 'reused')
        self.assertIn('SWIFT-38', Path(first['markdown']).read_text(encoding='utf-8'))
        self.assertTrue(first['needs_review'])

    def test_new_formats_keep_cell_slide_and_notebook_locations(self):
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        workbook = self.root / 'limits.xlsx'
        w = Workbook(); w.active.title = '预算'; w.active['B3'] = 'Budget: 28 hours'; w.save(workbook)
        slides = self.root / 'brief.pptx'
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Release lane: copper'
        p.save(slides)
        notebook = self.root / 'recipe.ipynb'
        notebook.write_text(json.dumps({'cells': [{'cell_type':'markdown','source':['Fold seed: 619']}]}), encoding='utf-8')
        for source, location, text in [(workbook, 'sheet', 'Budget: 28 hours'),
                                       (slides, 'slide', 'Release lane: copper'),
                                       (notebook, 'cell', 'Fold seed: 619')]:
            with self.subTest(source=source.name):
                result = cache_document(source, self.cache)
                body = Path(result['markdown']).read_text(encoding='utf-8')
                self.assertIn(text, body)
                self.assertIn(location, body)
                self.assertEqual(cache_document(source, self.cache)['status'], 'reused')

    def test_spreadsheet_row_keeps_labels_with_their_values(self):
        from openpyxl import Workbook
        source = self.root/'matrix.xlsx'
        book = Workbook(); book.active.append(['GPU budget', 28]); book.save(source)
        result = cache_document(source, self.cache)
        body = Path(result['markdown']).read_text(encoding='utf-8')
        self.assertIn('GPU budget | 28', body)

    def test_unknown_binary_without_prior_extraction_fails_explicitly(self):
        source = self.root/'unknown.bin'
        source.write_bytes(b'no generic decoder')
        with self.assertRaisesRegex(ValueError, 'from-markdown'):
            cache_document(source, self.cache)

    def test_archive_manifest_does_not_execute_or_extract_members(self):
        import zipfile
        source = self.root/'unsafe.zip'
        with zipfile.ZipFile(source,'w') as archive:
            archive.writestr('../UNREQUESTED_ACTION.txt','not executable data')
        result = cache_document(source,self.cache)
        self.assertIn('UNREQUESTED_ACTION',Path(result['markdown']).read_text(encoding='utf-8'))
        self.assertTrue(result['needs_review'])
        self.assertFalse((self.root/'UNREQUESTED_ACTION.txt').exists())


if __name__ == '__main__':
    unittest.main()
