"""Behavior tests for the document-cache CLI; fixtures stay in task work/."""
import json
import hashlib
import os
import re
from pathlib import Path
import subprocess
import sys
import unittest
import uuid

from docx import Document
from reportlab.pdfgen import canvas


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "document_cache.py"
WORK = Path.cwd() / "work" / "document-cache-tests"


class DocumentCacheTests(unittest.TestCase):
    def setUp(self):
        self.root = WORK / uuid.uuid4().hex
        self.root.mkdir(parents=True)
        self.cache = self.root / "cache"

    def make_word(self, path=None, text="数据与实验要求"):
        path = path or self.root / "规则.docx"
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading("实验要求", level=1)
        doc.add_paragraph(text)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "ID"
        table.cell(0, 1).text = "Score"
        table.cell(1, 0).text = "sample-1"
        table.cell(1, 1).text = "0.75"
        doc.add_paragraph("表格之后的说明")
        doc.save(path)
        return path

    def make_pdf(self, blank_second=False):
        path = self.root / "paper.pdf"
        pdf = canvas.Canvas(str(path))
        pdf.drawString(72, 720, "Baseline score: 0.75")
        pdf.showPage()
        if not blank_second:
            pdf.drawString(72, 720, "Experiment budget: 3 runs")
        pdf.showPage()
        pdf.save()
        return path

    def call(self, source, *extra, success=True):
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(source), "--cache-dir", str(self.cache), *extra],
            capture_output=True, text=True, encoding="utf-8",
        )
        if success:
            self.assertEqual(result.returncode, 0, result.stderr)
        else:
            self.assertNotEqual(result.returncode, 0)
        return json.loads(result.stdout)

    def test_word_keeps_unicode_table_order_and_block_locations(self):
        result = self.call(self.make_word())
        text = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertLess(text.index("数据与实验要求"), text.index("sample-1"))
        self.assertLess(text.index("sample-1"), text.index("表格之后的说明"))
        self.assertIn("| ID | Score |", text)
        self.assertIn("<!-- block 0003 -->", text)
        self.assertIn("# 实验要求", text)

    def test_second_request_reuses_cache_without_rewriting_files(self):
        source = self.make_word()
        first = self.call(source)
        md = Path(first["markdown"])
        meta = Path(first["metadata"])
        stamps = (md.stat().st_mtime_ns, meta.stat().st_mtime_ns)
        second = self.call(source)
        self.assertEqual(second["status"], "reused")
        self.assertEqual((md.stat().st_mtime_ns, meta.stat().st_mtime_ns), stamps)

    def test_source_update_rebuilds_cache_and_changes_extracted_content(self):
        source = self.make_word(text="原始要求")
        first = self.call(source)
        stamp = source.stat().st_mtime_ns
        self.make_word(source, "更新后的实际要求")
        os.utime(source, ns=(stamp + 2_000_000_000, stamp + 2_000_000_000))
        second = self.call(source)
        self.assertEqual(second["status"], "created")
        self.assertEqual(second["markdown"], first["markdown"])
        self.assertIn("更新后的实际要求", Path(second["markdown"]).read_text(encoding="utf-8"))

    def test_missing_cached_markdown_is_recreated(self):
        source = self.make_word()
        first = self.call(source)
        Path(first["markdown"]).unlink()
        second = self.call(source)
        self.assertEqual(second["status"], "created")
        self.assertIn("sample-1", Path(second["markdown"]).read_text(encoding="utf-8"))

    def test_same_filename_in_different_directories_does_not_share_cache(self):
        left = self.make_word(self.root / "left" / "same.docx", "左侧资料")
        right = self.make_word(self.root / "right" / "same.docx", "右侧资料")
        a, b = self.call(left), self.call(right)
        self.assertNotEqual(a["markdown"], b["markdown"])
        self.assertIn("左侧资料", Path(a["markdown"]).read_text(encoding="utf-8"))
        self.assertIn("右侧资料", Path(b["markdown"]).read_text(encoding="utf-8"))

    def test_pdf_retains_text_and_verified_physical_page_locations(self):
        result = self.call(self.make_pdf())
        text = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("## PDF page 1", text)
        self.assertIn("## PDF page 2", text)
        self.assertIn("Baseline score: 0.75", text)
        self.assertIn("Experiment budget: 3 runs", text)

    def test_pdf_without_text_on_a_page_flags_review_instead_of_complete_read(self):
        result = self.call(self.make_pdf(blank_second=True))
        meta = json.loads(Path(result["metadata"]).read_text(encoding="utf-8"))
        self.assertTrue(result["needs_review"])
        self.assertIn("PDF page 2", meta["review_locations"])
        self.assertIn("No text extracted", Path(result["markdown"]).read_text(encoding="utf-8"))

    def test_corrupt_source_does_not_replace_previous_good_cache_with_success(self):
        source = self.make_word()
        first = self.call(source)
        previous = Path(first["markdown"]).read_bytes()
        source.write_bytes(b"not a valid Word document")
        error = self.call(source, success=False)
        self.assertEqual(error["status"], "error")
        self.assertEqual(Path(first["markdown"]).read_bytes(), previous)

    def test_existing_format_skill_extraction_can_be_adopted_without_converting_again(self):
        source = self.make_pdf()
        existing = self.root / "already-read.md"
        existing.write_text("## PDF page 1\n已完成的正文提取，不再重复解析。", encoding="utf-8")
        result = self.call(source, "--from-markdown", str(existing))
        text = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("已完成的正文提取", text)
        self.assertTrue(result["needs_review"])
        self.assertEqual(self.call(source)["status"], "reused")

    def test_explicit_refresh_rebuilds_even_when_stat_is_unchanged(self):
        source = self.make_word()
        self.call(source)
        self.assertEqual(self.call(source, "--refresh")["status"], "created")

    def test_empty_word_body_is_flagged_for_review(self):
        source = self.root / "empty.docx"
        doc = Document()
        doc.add_paragraph("")
        doc.save(source)
        result = self.call(source)
        self.assertTrue(result["needs_review"])
        meta = json.loads(Path(result["metadata"]).read_text(encoding="utf-8"))
        self.assertIn("DOCX empty body extraction", meta["review_locations"])

    def test_word_paragraph_citations_do_not_count_tables_as_paragraphs(self):
        result = self.call(self.make_word())
        text = Path(result["markdown"]).read_text(encoding="utf-8")
        metadata = json.loads(Path(result["metadata"]).read_text(encoding="utf-8"))
        self.assertEqual(metadata["structure"], {"top_level_paragraphs": 3, "top_level_tables": 1})
        self.assertIn("<!-- table 0001 -->", text)
        self.assertEqual(text.count("<!-- paragraph "), 3)
        final_paragraph = text.split("<!-- paragraph 0003 -->", 1)[1]
        self.assertIn("表格之后的说明", final_paragraph)
        self.assertNotIn("sample-1", final_paragraph)

    def test_empty_paragraph_still_has_a_citable_index(self):
        source = self.root / "empty-between.docx"
        doc = Document()
        doc.add_paragraph("title", style="Title")
        doc.add_paragraph("")
        doc.add_paragraph("after empty paragraph")
        doc.save(source)
        result = self.call(source)
        text = Path(result["markdown"]).read_text(encoding="utf-8")
        self.assertIn("<!-- paragraph 0002 -->", text)
        self.assertIn("after empty paragraph", text.split("<!-- paragraph 0003 -->", 1)[1])

    def test_legacy_block_cache_is_reused_without_forced_upgrade(self):
        source = self.make_word()
        first = self.call(source)
        md, meta = Path(first["markdown"]), Path(first["metadata"])
        legacy = re.sub(r"<!-- (paragraph|table) \d+ -->\n\n", "", md.read_text(encoding="utf-8"))
        md.write_text(legacy, encoding="utf-8")
        record = json.loads(meta.read_text(encoding="utf-8"))
        record.pop("structure", None)
        record["output"] = {"size": md.stat().st_size, "mtime_ns": md.stat().st_mtime_ns,
                            "sha256": hashlib.sha256(md.read_bytes()).hexdigest()}
        meta.write_text(json.dumps(record), encoding="utf-8")
        stamps = (md.stat().st_mtime_ns, meta.stat().st_mtime_ns)
        second = self.call(source)
        self.assertEqual(second["status"], "reused")
        self.assertEqual((md.stat().st_mtime_ns, meta.stat().st_mtime_ns), stamps)
        self.assertEqual(md.read_text(encoding="utf-8"), legacy)


if __name__ == "__main__":
    unittest.main()
